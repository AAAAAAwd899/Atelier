"""Current integration/import/masking cases; native Anki is explicitly doubled."""
from __future__ import annotations
from pathlib import Path
from io import BytesIO
from copy import deepcopy
from types import SimpleNamespace as NS,ModuleType
import sys,zipfile,xml.etree.ElementTree as ET,json,importlib
import pytest
from addon import core,collection,documents,masks,importers
from addon.storage import Store
from tests.fakes import FakeCollection,AddNoteRequest,OpChanges
from tests.authoring_fixtures import png,presentation,enable_native_io

@pytest.fixture
def setup(tmp_path,monkeypatch):
    ac=ModuleType('anki.collection');ac.AddNoteRequest=AddNoteRequest;ac.OpChanges=OpChanges
    monkeypatch.setitem(sys.modules,'anki.collection',ac)
    s=Store(tmp_path);p=s.project({'title':'Source laboratory'});c=enable_native_io(FakeCollection(tmp_path))
    yield s,p,c
    c.close()

def mask(i=0,group=1,**kw):return {'id':f'{i:032x}','group':group,'x':.1,'y':.1,'w':.2,'h':.1,**kw}
def masked(s,p,**kw):
    blob=s.save_blob(png(),'png')
    return s.save_draft({'project_id':p['id'],'kind':'occlusion','front':'Recall the hidden labels.','back':'Check the definitions in their source.','deck_id':2,
      'media':[{'blob':blob,'side':'front'}],'masks':[mask(1),mask(2,1,x=.4),mask(3,2,y=.4)],'hide_others':True,**kw})

@pytest.mark.parametrize('field,value',[(k,v) for k in ('x','y','w','h') for v in (None,True,float('nan'),float('inf'),-1,1.2)])
def test_bad_mask_coordinates(field,value):
    with pytest.raises(core.ValidationError):masks.validate_masks([mask(**{field:value})])
@pytest.mark.parametrize('group',[0,501,True,'nan',1.5])
def test_bad_mask_groups(group):
    with pytest.raises(core.ValidationError):masks.validate_masks([mask(group=group)])
def test_groups_same_card_and_hide_policy():
    data=[mask(1),mask(2),mask(3,2)]
    assert masks.question_count(data)==2
    value=masks.serialize(data)
    assert value.count('{{c1::image-occlusion:rect:')==2 and ':oi=1' in value
    assert ':oi=' not in masks.serialize(data,False)
def test_masks_bounded_and_unique():
    for raw in ([mask(1),mask(1)],[mask(w=.95)], [mask(h=.00001)], [mask(i) for i in range(1001)]):
        with pytest.raises(core.ValidationError):masks.validate_masks(raw)
def test_zero_mask_local_draft_but_no_publication(setup):
    s,p,c=setup;d=masked(s,p,masks=[])
    with pytest.raises(core.ValidationError):collection.preflight(c,s,[d['id']])
    assert not c.writes
def test_native_grouped_publication_one_note_two_cards(setup):
    s,p,c=setup;d=masked(s,p);pf=collection.preflight(c,s,[d['id']])
    assert pf['total_cards']==2 and not c.models.items
    r=collection.commit(c,s,pf,set());n=c.get_note(r.data['note_ids'][0])
    assert n.note_type()['originalStockKind']==6 and n.fields[0].count('{{c1::')==2
    assert len(c.card_ids_of_note(n.id))==2
    assert c.db.list('select did from cards')==[2,2] and c.db.all('select * from revlog')==[]
    assert c.writes==[('add_notes',1)]
    c.undo();assert c.note_count()==0

def test_no_duplicate_republish_on_retry(setup):
    s,p,c=setup;d=masked(s,p);r=collection.commit(c,s,collection.preflight(c,s,[d['id']]),set());
    with pytest.raises(core.ValidationError):collection.preflight(c,s,[d['id']])
    assert c.note_count()==1 and len(c.db.all('select * from cards'))==2

def test_mask_source_only_raster_and_native_type(setup):
    s,p,c=setup;d=masked(s,p);d['media'].append({'blob':s.save_blob(png(),'png'),'side':'front'})
    d=s.save_draft(d)
    with pytest.raises(core.ValidationError):collection.preflight(c,s,[d['id']])

def test_native_model_missing_no_invented_template(setup):
    s,p,c=setup;del c._backend;d=masked(s,p)
    with pytest.raises(core.ValidationError):collection.commit(c,s,collection.preflight(c,s,[d['id']]),set())
    assert not c.notes

def test_native_metadata_invalid_does_not_create_notes(setup):
    s,p,c=setup;c._backend.get_image_occlusion_fields=lambda **kw:NS(occlusions=0,image=0,header=0,back_extra=0)
    d=masked(s,p)
    with pytest.raises(core.ValidationError):collection.commit(c,s,collection.preflight(c,s,[d['id']]),set())
    assert not c.notes

def test_existing_card_schedules_unchanged_on_mask_add(setup):
    s,p,c=setup;n=c.seed('Existing mature card','Answer');c.db.conn.execute('update cards set queue=2,type=2,ivl=91,reps=23')
    before=c.db.all('select * from cards where nid=?',n.id);d=masked(s,p)
    collection.commit(c,s,collection.preflight(c,s,[d['id']]),set())
    assert before==c.db.all('select * from cards where nid=?',n.id)

def test_pptx_real_order_notes_table_figures_hidden(setup,tmp_path):
    s,p,c=setup;result=documents.capture_file(s,p['id'],presentation(tmp_path/'lecture.pptx'))
    pages=s.get('version',result['source']['version_id'])['pages']
    assert 'Lesson 2' in pages[0]['text'] and 'Lesson 1' in pages[1]['text']
    assert 'Item 2 | A verified explanation' in pages[0]['text']
    assert 'Speaker detail for lesson 2' in pages[0]['notes']
    assert pages[0]['hidden_slide'] and not pages[1]['hidden_slide']
    assert pages[0]['assets'][0]['blob']==pages[1]['assets'][0]['blob']
    result=documents.render_page(s,result['source']['version_id'],1)
    assert not result['image'] and result['assets'][0]['image'].startswith('data:image/png')
    assert 'not' in s.get('version',s.list('source')[0]['version_id'])['metadata']['extraction_note'].lower()

@pytest.mark.parametrize('target',['../../escape.png','https://remote.test/x','..\\evil','file:///x','a%2Fb','a#b','a?b'])
def test_archive_relationships_not_paths_or_network(target):
    with pytest.raises(core.ValidationError):importers.resolve_part('ppt/presentation.xml',target)
def test_internal_image_relationship_valid():assert importers.resolve_part('ppt/slides/slide1.xml','../media/image1.png')=='ppt/media/image1.png'
def test_external_relationships_not_downloaded():
    b=BytesIO()
    with zipfile.ZipFile(b,'w') as z:z.writestr('ppt/_rels/presentation.xml.rels','<Relationships><Relationship Id="r1" Target="https://example.test/a.png" TargetMode="External" Type="image"/></Relationships>')
    with zipfile.ZipFile(b) as z:assert importers.relationships(z,'ppt/presentation.xml',documents._xml)=={}
@pytest.mark.parametrize('enc',['utf-8','utf-16','utf-16-be'])
def test_xml_entities_refused(enc):
    with pytest.raises(core.ValidationError):documents._xml('<!DOCTYPE a [<!ENTITY x SYSTEM "file:///etc/passwd">]><a>&x;</a>'.encode(enc))
@pytest.mark.parametrize('ext,raw,expected', [('csv','Term,Meaning\n"one, two",Comparison\n','one, two'),('tsv','Term\tMeaning\nx\tValue\n','x | Value')])
def test_delimited_sources(setup,tmp_path,ext,raw,expected):
    s,p,c=setup;path=tmp_path/('source.'+ext);path.write_text(raw)
    r=documents.capture_file(s,p['id'],path);v=s.get('version',r['source']['version_id']);assert expected in v['pages'][0]['text'] and not c.writes

def test_docx_tables_and_figures(setup,tmp_path):
    from docx import Document
    d=Document();d.add_heading('Original notes',0);tb=d.add_table(rows=1,cols=2);tb.cell(0,0).text='Question';tb.cell(0,1).text='Answer';d.add_picture(BytesIO(png()));path=tmp_path/'notes.docx';d.save(path)
    s,p,c=setup;r=documents.capture_file(s,p['id'],path);v=s.get('version',r['source']['version_id']);assert 'Question | Answer' in v['pages'][0]['text'] and len(v['pages'][0]['assets'])==1

@pytest.mark.parametrize('ext',['odt','odp'])
def test_opendocument_sources(setup,tmp_path,ext):
    s,p,c=setup;path=tmp_path/('notes.'+ext)
    xml='<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"><office:body><draw:page><text:p>Original section</text:p></draw:page></office:body></office:document-content>'
    with zipfile.ZipFile(path,'w') as z:z.writestr('content.xml',xml)
    r=documents.capture_file(s,p['id'],path);assert 'Original section' in s.get('version',r['source']['version_id'])['pages'][0]['text']

def make_pdf(path,n):
    from reportlab.pdfgen.canvas import Canvas
    p=Canvas(str(path))
    for i in range(n):p.drawString(50,750,f'Slide {i+1}');p.showPage()
    p.save();return path

def test_pdf_layout_attachment_preserves_source_and_checks_count(setup,tmp_path):
    s,p,c=setup;r=documents.capture_file(s,p['id'],presentation(tmp_path/'lecture.pptx'));vid=r['source']['version_id'];v=s.get('version',vid)
    with pytest.raises(core.ValidationError):documents.attach_pdf_layout(s,vid,make_pdf(tmp_path/'bad.pdf',1))
    assert s.get('version',vid)==v
    documents.attach_pdf_layout(s,vid,make_pdf(tmp_path/'good.pdf',2));after=s.get('version',vid)
    assert after['blob']==v['blob'] and after['pages']==v['pages'] and after['metadata']['visual_pdf'].endswith('.pdf')
    assert after['revision']==v['revision']+1

def test_source_assets_roundtrip_and_id_detach(setup,tmp_path):
    s,p,c=setup;r=documents.capture_file(s,p['id'],presentation(tmp_path/'lecture.pptx'));documents.attach_pdf_layout(s,r['source']['version_id'],make_pdf(tmp_path/'matching.pdf',2))
    d=masked(s,p);path=tmp_path/'library.zip';documents.export_workspace(s,path)
    target=Store(tmp_path/'new');documents.import_workspace(target,path);v=target.list('version')[0]
    assert 'Speaker detail' in v['pages'][0]['notes'] and v['pages'][0]['assets']
    assert target.blob_path(v['metadata']['visual_pdf']).exists()
    assert target.list('draft')[0]['masks']==d['masks'] and target.list('draft')[0]['deck_id']==1

def test_fresh_suite_theme_only_fresh_store(tmp_path):
    s=Store(tmp_path,initial_theme='atelier');assert s.settings()['theme']=='atelier'
    s.set_settings({'theme':'paper'});before=s.settings();Store(tmp_path,initial_theme='atelier');assert s.settings()==before
